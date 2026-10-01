import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
BG = RGBColor(11, 18, 32)        # #0B1220
CARD = RGBColor(20, 30, 51)      # #141E33
CIRC = RGBColor(30, 43, 72)      # #1E2B48
AMBER = RGBColor(245, 165, 36)   # #F5A524
WHITE = RGBColor(241, 244, 250)  # #F1F4FA
MUTED = RGBColor(154, 169, 194)  # #9AA9C2
GREEN = RGBColor(61, 220, 151)   # #3DDC97
RED = RGBColor(255, 92, 92)      # #FF5C5C

HEAD = 'Cambria'
BODY = 'Calibri'

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SHOT1 = os.path.join(SCRIPT_DIR, 'screenshots', '1_crowd_risk_ai_meter.png')
SHOT2 = os.path.join(SCRIPT_DIR, 'screenshots', '2_sos_and_alerts.png')
SHOT3 = os.path.join(SCRIPT_DIR, 'screenshots', '3_mesh_relay_network.png')

blank_layout = prs.slide_layouts[6]

def set_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG

def add_header(slide, title_text, width=12.0):
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(width), Inches(0.9))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = HEAD
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Accent line under header
    line = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(0.8), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = AMBER
    line.line.color.rgb = AMBER

# ==========================================
# SLIDE 1: Title Slide
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_background(s1)

tb = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(10.0), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = "iQOO HACKATHON 2026  ·  GRAND FINALE APPLICATION"
p.font.name = BODY
p.font.size = Pt(13)
p.font.color.rgb = MUTED

tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(10.0), Inches(1.3))
p = tb.text_frame.paragraphs[0]
p.text = "Festival Guardian"
p.font.name = HEAD
p.font.size = Pt(56)
p.font.bold = True
p.font.color.rgb = WHITE

tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(10.0), Inches(0.7))
p = tb.text_frame.paragraphs[0]
p.text = "Predict. Respond. Relay."
p.font.name = HEAD
p.font.size = Pt(26)
p.font.italic = True
p.font.color.rgb = AMBER

tb = s1.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(9.0), Inches(1.1))
p = tb.text_frame.paragraphs[0]
p.text = "Point your phone at a crowd and get a live stampede-risk score. If the network jams, alerts hop from phone to phone."
p.font.name = BODY
p.font.size = Pt(17)
p.font.color.rgb = MUTED

# Divider
div = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.0), Inches(9.5), Inches(0.03))
div.fill.solid()
div.fill.fore_color.rgb = CARD
div.line.color.rgb = CARD

tb = s1.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(9.0), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = "TEAM FALLING STARS"
p.font.name = BODY
p.font.size = Pt(13)
p.font.color.rgb = AMBER
p.font.bold = True

tb = s1.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(9.0), Inches(0.5))
p = tb.text_frame.paragraphs[0]
p.text = "Poornachandra P M   ·   Dushyanth M   ·   M Sujay"
p.font.name = BODY
p.font.size = Pt(16)
p.font.color.rgb = WHITE

# ==========================================
# SLIDE 2: Problem
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_background(s2)
add_header(s2, "Crowds turn deadly in minutes")

probs = [
    ("Nobody sees it building", "Density rises gradually. By the time it feels dangerous, exits are already blocked."),
    ("Networks fail at peak", "Thousands of phones in one spot overload towers, so calls and messages fail."),
    ("Help arrives late", "Organizers get no early signal and no location, so response starts after the crush.")
]

for i, (head_t, body_t) in enumerate(probs):
    x = 0.8 + i * 4.0
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(3.7), Inches(4.0))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = CARD
    
    # Text
    tb = s2.shapes.add_textbox(Inches(x + 0.35), Inches(2.4), Inches(3.0), Inches(0.8))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = head_t
    p.font.name = BODY
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    tb2 = s2.shapes.add_textbox(Inches(x + 0.35), Inches(3.4), Inches(3.0), Inches(2.2))
    tb2.text_frame.word_wrap = True
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = body_t
    p2.font.name = BODY
    p2.font.size = Pt(15)
    p2.font.color.rgb = MUTED

tb = s2.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = "Festival grounds, temple queues, railway platforms."
p.font.name = BODY
p.font.size = Pt(14)
p.font.italic = True
p.font.color.rgb = AMBER

# ==========================================
# SLIDE 3: Solution / Live Demo
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_background(s3)
add_header(s3, "Point. Read. Warn.")

rows = [
    ("Point the camera", "Detection runs on the phone's NPU/GPU, with no upload."),
    ("Live risk meter", "Density and flow combine into one clear 0–100 risk score."),
    ("Early warning", "Organizers are alerted before a crush or stampede forms.")
]

for i, (title_t, desc_t) in enumerate(rows):
    y = 2.0 + i * 1.5
    tb = s3.shapes.add_textbox(Inches(0.8), Inches(y), Inches(6.8), Inches(1.3))
    tb.text_frame.word_wrap = True
    p1 = tb.text_frame.paragraphs[0]
    p1.text = title_t
    p1.font.name = BODY
    p1.font.size = Pt(21)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p2 = tb.text_frame.add_paragraph()
    p2.text = desc_t
    p2.font.name = BODY
    p2.font.size = Pt(15)
    p2.font.color.rgb = MUTED

# Screenshot on Right
if os.path.exists(SHOT1):
    frame = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(0.8), Inches(3.8), Inches(5.9))
    frame.fill.solid()
    frame.fill.fore_color.rgb = CARD
    frame.line.color.rgb = AMBER
    frame.line.width = Pt(2)
    s3.shapes.add_picture(SHOT1, Inches(8.6), Inches(0.9), Inches(3.6), Inches(5.7))

tb = s3.shapes.add_textbox(Inches(8.0), Inches(6.8), Inches(4.8), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = "Actual screenshot — live prototype (demo/test data)"
p.font.name = BODY
p.font.size = Pt(11)
p.font.italic = True
p.font.color.rgb = MUTED
p.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 4: How It Works
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_background(s4)
add_header(s4, "How it works, today vs. the Finale build")

steps = [
    ("1", "Camera", "Live feed from the phone"),
    ("2", "Detect", "TensorFlow.js (COCO-SSD) runs on-device, in-browser, over WebGL"),
    ("3", "Score", "Density + flow become a 0-100 risk score"),
    ("4", "Alert", "A typed Alert object is created, with TTL + dedup"),
    ("5", "Relay", "Working today same-device; native phone-to-phone at the Finale")
]

for i, (num, name, desc) in enumerate(steps):
    x = 0.8 + i * 2.38
    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(2.2), Inches(4.1))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = CARD
    
    tb = s4.shapes.add_textbox(Inches(x + 0.15), Inches(2.1), Inches(1.9), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = num
    p.font.name = HEAD
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = AMBER
    
    tb2 = s4.shapes.add_textbox(Inches(x + 0.15), Inches(2.8), Inches(1.9), Inches(0.6))
    tb2.text_frame.word_wrap = True
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = name
    p2.font.name = BODY
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER
    
    tb3 = s4.shapes.add_textbox(Inches(x + 0.15), Inches(3.6), Inches(1.9), Inches(2.2))
    tb3.text_frame.word_wrap = True
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = desc
    p3.font.name = BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = MUTED
    p3.alignment = PP_ALIGN.CENTER

tb = s4.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(12.0), Inches(0.5))
p = tb.text_frame.paragraphs[0]
p.text = "Today: real on-device detection, in a web prototype. Finale: ported to native Android + Nearby Connections for true phone-to-phone offline relay."
p.font.name = BODY
p.font.size = Pt(13)
p.font.italic = True
p.font.color.rgb = AMBER

# ==========================================
# SLIDE 5: Offline SOS Relay
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_background(s5)
add_header(s5, "Offline SOS relay — the Finale target")

tb = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.8), Inches(0.8))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = "No network? Alerts hop from phone to phone until one reaches a connection. The relay logic (TTL, dedup, typed alerts) is built; the phone-to-phone hop is the Finale build."
p.font.name = BODY
p.font.size = Pt(16)
p.font.color.rgb = MUTED

nodes = [
    ("Person in the crowd", AMBER, True),
    ("Nearby phone (Hop 1)", CARD, False),
    ("Nearby phone (Hop 2)", CARD, False),
    ("Volunteer with signal", CARD, False)
]

for i, (name, bg_col, is_first) in enumerate(nodes):
    x = 0.8 + i * 2.95
    node_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.2), Inches(2.5), Inches(2.2))
    node_box.fill.solid()
    node_box.fill.fore_color.rgb = bg_col
    node_box.line.color.rgb = AMBER
    node_box.line.width = Pt(1.5)
    
    tb = s5.shapes.add_textbox(Inches(x + 0.1), Inches(3.8), Inches(2.3), Inches(1.0))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = name
    p.font.name = BODY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BG if is_first else WHITE
    p.alignment = PP_ALIGN.CENTER
    
    if i < 3:
        arr = s5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.55), Inches(4.1), Inches(0.35), Inches(0.3))
        arr.fill.solid()
        arr.fill.fore_color.rgb = AMBER
        arr.line.color.rgb = AMBER

tb = s5.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(12.0), Inches(0.5))
p = tb.text_frame.paragraphs[0]
p.text = "Today: relay engine verified across browser tabs (same device). Finale: Android Nearby Connections, two real phones, one offline."
p.font.name = BODY
p.font.size = Pt(13)
p.font.italic = True
p.font.color.rgb = AMBER

# ==========================================
# SLIDE 6: Who Gets the Alert
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_background(s6)
add_header(s6, "Who gets the alert")

who = [
    ("Event organizers", "Live risk score, location and photo on their phone."),
    ("Emergency services (108 / 112)", "A pre-filled emergency call with the exact location and incident type."),
    ("Family and nearby volunteers", "Trusted contacts and nearby app users get the alert pin and guidance.")
]

for i, (title_t, desc_t) in enumerate(who):
    y = 2.0 + i * 1.4
    card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(11.7), Inches(1.15))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = CARD
    
    tb = s6.shapes.add_textbox(Inches(1.2), Inches(y + 0.15), Inches(11.0), Inches(0.8))
    tb.text_frame.word_wrap = True
    p1 = tb.text_frame.paragraphs[0]
    p1.text = title_t
    p1.font.name = BODY
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p2 = tb.text_frame.add_paragraph()
    p2.text = desc_t
    p2.font.name = BODY
    p2.font.size = Pt(14)
    p2.font.color.rgb = MUTED

# ==========================================
# SLIDE 7: Phone-first — Today vs. Finale
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_background(s7)
add_header(s7, "Today vs. the Finale build")

cols = [
    ("Built today (web prototype)", [
        "Real on-device detection: TF.js + COCO-SSD over WebGL",
        "Alert engine: typed alerts, TTL, dedup — all working",
        "Relay verified same-device (browser tabs)"
    ]),
    ("Finale: Red Light (phone only)", [
        "Port detection to TFLite, run on the Snapdragon NPU",
        "Android Nearby Connections for real phone-to-phone hops",
        "Live demo: 2-3 iQOO phones, one offline"
    ])
]

for i, (col_head, items) in enumerate(cols):
    x = 0.8 + i * 6.0
    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(5.6), Inches(4.2))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = CARD
    
    tb = s7.shapes.add_textbox(Inches(x + 0.4), Inches(2.3), Inches(4.8), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    p.text = col_head
    p.font.name = BODY
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    tb2 = s7.shapes.add_textbox(Inches(x + 0.4), Inches(3.2), Inches(4.8), Inches(2.6))
    tb2.text_frame.word_wrap = True
    tf2 = tb2.text_frame
    for j, itm in enumerate(items):
        p_item = tf2.paragraphs[0] if j == 0 else tf2.add_paragraph()
        p_item.text = f"•  {itm}"
        p_item.font.name = BODY
        p_item.font.size = Pt(15)
        p_item.font.color.rgb = MUTED
        p_item.space_after = Pt(14)

tb = s7.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(12.0), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = "Scored on: phone-first execution, AI integration, Office Kit, real-world relevance, pitch."
p.font.name = BODY
p.font.size = Pt(13)
p.font.italic = True
p.font.color.rgb = AMBER

# ==========================================
# SLIDE 8: Live Prototype — Not a Mockup
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
set_background(s8)
add_header(s8, "Live prototype — not a mockup")

tb = s8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.8), Inches(0.7))
tb.text_frame.word_wrap = True
p = tb.text_frame.paragraphs[0]
p.text = "Actual screenshots from the working web prototype. Names, coordinates and counts are demo/test data. Mesh peers shown are seeded for the UI; live device-to-device relay is the Finale build."
p.font.name = BODY
p.font.size = Pt(14)
p.font.color.rgb = MUTED

shots = [
    (SHOT2, "SOS, theft & volunteer alerts"),
    (SHOT3, "Mesh network — connected peers")
]

for i, (path, label) in enumerate(shots):
    x = 2.4 + i * 4.6
    if os.path.exists(path):
        frame = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.4), Inches(3.8), Inches(4.2))
        frame.fill.solid()
        frame.fill.fore_color.rgb = CARD
        frame.line.color.rgb = AMBER
        frame.line.width = Pt(1.5)
        s8.shapes.add_picture(path, Inches(x + 0.1), Inches(2.5), Inches(3.6), Inches(4.0))
    
    tb = s8.shapes.add_textbox(Inches(x), Inches(6.7), Inches(3.8), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = label
    p.font.name = BODY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 9: 48-Hour Plan & Thank You
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
set_background(s9)
add_header(s9, "48 hours: what we ship")

plan = [
    ("0-12h", "Detection + meter", "On-device counting and a live risk score"),
    ("12-30h", "Offline relay", "Two-phone SOS hop working end to end"),
    ("30-42h", "Alerts + first aid", "Location, photo and a voice-guided first-aid card"),
    ("42-48h", "Polish + pitch", "Demo rehearsal, backup video, final deck")
]

for i, (hours, title_t, desc_t) in enumerate(plan):
    x = 0.8 + i * 2.95
    card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(2.7), Inches(3.0))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = CARD
    
    tb = s9.shapes.add_textbox(Inches(x + 0.2), Inches(2.2), Inches(2.3), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = hours
    p.font.name = HEAD
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = AMBER
    
    tb2 = s9.shapes.add_textbox(Inches(x + 0.2), Inches(2.8), Inches(2.3), Inches(0.5))
    tb2.text_frame.word_wrap = True
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = title_t
    p2.font.name = BODY
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    
    tb3 = s9.shapes.add_textbox(Inches(x + 0.2), Inches(3.5), Inches(2.3), Inches(1.3))
    tb3.text_frame.word_wrap = True
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = desc_t
    p3.font.name = BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = MUTED

tb = s9.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(12.0), Inches(0.5))
p = tb.text_frame.paragraphs[0]
p.text = "Festival Guardian: predict, respond, relay."
p.font.name = HEAD
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = MUTED

tb_ty = s9.shapes.add_textbox(Inches(0.8), Inches(5.9), Inches(12.0), Inches(1.0))
p_ty = tb_ty.text_frame.paragraphs[0]
p_ty.text = "Thank You"
p_ty.font.name = HEAD
p_ty.font.size = Pt(44)
p_ty.font.bold = True
p_ty.font.color.rgb = AMBER

# Save presentation
out_path = os.path.join(SCRIPT_DIR, 'Festival-Guardian-Pitch-Deck.pptx')
prs.save(out_path)
print(f"Deck saved successfully at: {out_path}")
